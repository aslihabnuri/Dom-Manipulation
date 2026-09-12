# -*- coding: utf-8 -*-
"""
Naskah presentasi SLIDE 1-7 (urutan terbaru, 22 slide) untuk pembicara pertama: Aulia.
Deck: "Inventory Management sebagai Keunggulan Kompetitif Amazon" - Kelompok 4.
Urutan pembicara keseluruhan: Aulia, Fitra, Bagaskoro.
Jalankan:  python3 build_naskah_aulia.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Naskah_Presentasi_Slide_1-7_Aulia_Kelompok_4.docx")
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5F, 0x5F, 0x5F)
MAGENTA = RGBColor(0xC4, 0x18, 0x5C)
PINK = "F6D5DA"

# (nomor, judul, durasi, inti, di layar, naskah[], transisi, tips[])
SLIDES = [
    (1, "Sampul", "40 detik",
     "Perkenalan kelompok dan janji isi presentasi.",
     "Judul besar, logo Amazon, dan pita nama Kelompok 4 di bagian bawah.",
     ["Selamat pagi Ibu Prof. Nur Aini, dan teman-teman semua.",
      "Kami dari Kelompok 4. Saya Aulia Sisca Rahmadiyanti, bersama Fitra Aidila dan Bagaskoro.",
      "Studi kasus yang kami bahas hari ini berjudul Inventory Management sebagai Keunggulan Kompetitif Amazon.",
      "Biasanya persediaan dianggap beban. Barang menumpuk di gudang, dan uang perusahaan tertahan di situ. "
      "Tetapi pada Amazon justru sebaliknya. Persediaan adalah alasan utama Amazon bisa lebih unggul dari "
      "pesaingnya. Itulah yang ingin kami tunjukkan pagi ini."],
     "Supaya mudah diikuti, kami mulai dari peta pembahasannya dulu.",
     ["Berdiri tegak dan jangan terburu-buru. Kalimat pertama pelan saja, audiens masih menyesuaikan diri.",
      "Sebut nama dua rekan Anda sambil menunjuk ke arah mereka, jangan membaca dari layar.",
      "Kalimat pembuka tentang persediaan yang dianggap beban adalah pengait perhatian. Ucapkan dengan jelas, "
      "lalu berhenti sejenak sebelum pindah slide."]),

    (2, "Tujuan dan Alur Pembahasan", "1 menit",
     "Peta jalan presentasi dalam enam langkah.",
     "Kotak tujuan di atas, lalu enam butir alur pembahasan dalam dua kolom.",
     ["Tujuan kami sederhana: melihat bagaimana Amazon mengelola persediaan secara menyeluruh.",
      "Ada tiga pertanyaan yang ingin kami jawab. Pertama, kenapa Amazon akhirnya memilih menyimpan stok "
      "sendiri. Kedua, bagaimana satu pesanan diproses sampai tiba di tangan pelanggan. Ketiga, pilihan sulit "
      "apa saja yang harus diambil Amazon di sepanjang jalan itu.",
      "Alurnya ada enam bagian. Kami mulai dari kerangka berpikir. Lalu kisah suksesnya, dari toko virtual "
      "menjadi pengelola persediaan. Ketiga, proses pemenuhan pesanan. Keempat, analisis trade-off. Kelima, "
      "perkembangan terbaru tahun 2023 dan 2025. Terakhir, lesson learned.",
      "Satu catatan tentang urutannya. Kami sengaja membahas strategi dulu, baru hal teknis. Alasannya, "
      "keputusan soal stok selalu mengikuti strategi yang sudah dipilih perusahaan, bukan sebaliknya."],
     "Sebelum masuk ke butir pertama, kami perkenalkan dulu perusahaan yang kami bahas.",
     ["Tunjuk keenam butir dengan cepat. Jangan dibacakan satu per satu, audiens bisa membaca sendiri.",
      "Kalimat terakhir tentang strategi dulu baru teknis adalah cara berpikir yang dinilai dosen. "
      "Jangan dilewati.",
      "Kalau waktu mepet, slide ini boleh dipersingkat menjadi dua kalimat: tujuan, lalu sebut enam bagiannya "
      "dalam satu tarikan napas."]),

    (3, "Tentang Amazon", "1 menit 15 detik",
     "Amazon itu perusahaan besar, tetapi intinya satu: fokus pada pelanggan.",
     "Profil Amazon: tahun berdiri, fokus awal, jangkauan global, karyawan, bisnis utama, nilai perusahaan, "
     "dan posisi di dunia.",
     ["Amazon adalah perusahaan teknologi dan e-commerce global. Didirikan tahun 1994 oleh Jeff Bezos di "
      "Seattle, Amerika Serikat.",
      "Awalnya Amazon hanya menjual buku secara online. Sekarang Amazon beroperasi di lebih dari 20 negara, "
      "melayani ratusan juta pelanggan, dengan lebih dari 1,5 juta karyawan di seluruh dunia.",
      "Bisnisnya juga tidak lagi hanya toko online. Ada Amazon Web Services untuk layanan cloud, ada Prime "
      "Video untuk streaming, dan ada jaringan logistiknya sendiri. Amazon kini termasuk perusahaan dengan "
      "nilai pasar terbesar di dunia.",
      "Tetapi dari semua itu, yang paling penting untuk pembahasan kita adalah nilai perusahaannya: "
      "Customer Obsession. Artinya, pelanggan menjadi titik awal setiap keputusan.",
      "Tolong diingat kalimat ini, karena akan kita pakai berkali-kali. Semua keputusan Amazon soal gudang, "
      "stok, dan robot, semuanya berangkat dari satu pertanyaan: apa yang paling baik untuk pelanggan?"],
     "Supaya tidak berhenti di angka, kami ingin teman-teman melihat sendiri seperti apa gudang Amazon itu.",
     ["Jangan membaca semua kotak di layar. Sebut yang penting saja, lalu tunjuk sisanya sekilas.",
      "Poin yang wajib diucapkan adalah Customer Obsession, karena itu benang merah seluruh presentasi.",
      "Kalau ada yang bertanya, tahun 1994 adalah tahun perusahaan didirikan. Situsnya baru mulai melayani "
      "pembeli tahun 1995, seperti yang muncul di slide kisah sukses nanti."]),

    (4, "Inventory Management Amazon (video)", "1 menit",
     "Gambaran nyata suasana gudang sebelum masuk teori.",
     "Cuplikan video Inside an Amazon Same-Day Center, dan foto gudang Amazon.",
     ["Ini cuplikan singkat dari dalam salah satu pusat pemenuhan pesanan Amazon.",
      "[Putar video sekitar 30 sampai 45 detik.]",
      "Ada tiga hal yang ingin kami minta teman-teman perhatikan dari video tadi.",
      "Pertama, barangnya sangat banyak dan sangat beragam. Kedua, orang dan barang bergerak cepat, hampir "
      "tidak ada yang berhenti. Ketiga, setiap barang punya label dan selalu dipindai.",
      "Banyak, cepat, dan akurat. Tiga kata itu adalah inti dari seluruh pembahasan kita hari ini."],
     "Sekarang kita mundur sebentar dari gudangnya, dan lihat cara berpikir di baliknya.",
     ["Siapkan videonya sebelum presentasi dimulai. Pastikan suara laptop menyala dan volumenya cukup.",
      "Jangan memutar video lebih dari 45 detik, audiens akan kehilangan fokus.",
      "Kalau video gagal diputar, jangan panik. Tunjuk fotonya, lalu langsung ke kalimat tentang tiga hal "
      "yang perlu diperhatikan.",
      "Catatan tampilan: judul slide ini masih tertutup video, kata Amazon terpotong. Kalau sempat, geser "
      "kotak judulnya ke atas sebelum presentasi."]),

    (5, "Kerangka Berpikir", "1 menit 15 detik",
     "Tiga pertanyaan dasar persediaan, dan yang ketiga paling menentukan bagi Amazon.",
     "Tiga pertanyaan: How much, When, Where; ditutup kotak tentang lokasi penyimpanan.",
     ["Sebelum bicara teknis, setiap perusahaan yang punya persediaan harus menjawab tiga pertanyaan dasar.",
      "How much. Berapa banyak stok yang perlu disiapkan untuk setiap produk?",
      "When. Kapan harus memesan lagi supaya stoknya tidak habis?",
      "Where. Di gudang mana setiap produk sebaiknya disimpan? Untuk Amazon ini berat, karena stoknya "
      "tersebar di lebih dari 150 gudang.",
      "Di kebanyakan perusahaan, dua pertanyaan pertama yang paling sering dibahas. Tetapi di Amazon, justru "
      "pertanyaan ketiga yang paling menentukan.",
      "Amazon menempatkan produk sedekat mungkin dengan pelanggan, bahkan sebelum pesanan masuk. Jadi ketika "
      "pelanggan menekan tombol beli, barangnya sudah berada di kota yang sama. Itulah sebabnya pesanan bisa "
      "diproses dan dikirim jauh lebih cepat."],
     "Setelah tahu tiga pertanyaannya, mari kita lihat struktur dari setiap keputusan persediaan.",
     ["Sebut ketiga istilah bahasa Inggrisnya persis seperti di layar, karena itu istilah baku di mata "
      "kuliah ini.",
      "Tekankan frasa sebelum pesanan masuk. Ini kunci yang akan muncul lagi di bagian proses dan "
      "perkembangan terbaru.",
      "Kalau ingin interaktif, boleh bertanya ke audiens: dari tiga pertanyaan ini, menurut teman-teman mana "
      "yang paling sulit untuk Amazon? Lalu lanjutkan ke jawabannya."]),

    (6, "Struktur Keputusan", "1 menit 30 detik",
     "Tujuan, keputusan, dan batasan; tingkat pelayanan itu batas, bukan tujuan.",
     "Tiga bagian: Objective, Decision Variables, Constraints; ditutup kotak tentang tingkat pelayanan.",
     ["Setiap keputusan soal persediaan, di perusahaan mana pun, selalu punya tiga bagian yang sama.",
      "Bagian pertama, objective atau tujuannya. Umumnya menekan tiga hal: biaya penyimpanan, biaya "
      "pemesanan, dan kerugian saat kehabisan stok. Amazon menambah satu lagi, yaitu biaya retur atau "
      "kesalahan pesanan. Bagi Amazon, salah kirim itu mahal sekali.",
      "Bagian kedua, decision variables, yaitu apa saja yang harus diputuskan. Berapa jumlah barang yang "
      "dipesan, kapan memesan kembali, dan berapa stok cadangannya. Amazon menambah satu keputusan lagi: "
      "produk apa disimpan, dan di gudang mana.",
      "Bagian ketiga, constraints atau batasannya. Ada target pelayanan, kapasitas gudang, dan waktu tunggu "
      "barang. Bagi Amazon batasannya berat: harga harus tetap kompetitif, pengiriman harus cepat, dan retur "
      "harus mudah.",
      "Ada satu hal yang ingin kami tekankan dari slide ini. Tingkat pelayanan itu bukan tujuan, melainkan "
      "batas yang harus dipenuhi. Jadi Amazon bukan mengejar layanan setinggi mungkin dengan biaya berapa "
      "pun. Amazon menekan biaya, dengan syarat janji kepada pelanggan tetap ditepati."],
     "Kalau tujuannya menekan biaya, muncul pertanyaan berikutnya: sebenarnya berapa jumlah stok yang pas?",
     ["Gunakan tiga jari untuk menandai tiga bagian, audiens akan lebih mudah mengikuti.",
      "Kalimat terakhir adalah pesan utama slide ini. Ucapkan pelan, lalu berhenti sejenak sebelum lanjut.",
      "Kalau ditanya beda tujuan dan batas: tujuan itu yang dikejar sampai sekecil mungkin, yaitu biaya. "
      "Batas itu yang tidak boleh dilanggar, yaitu janji layanan."]),

    (7, "Dilema Persediaan dan Aturan Keputusan", "1 menit 15 detik",
     "Stok banyak dan stok sedikit sama-sama merugikan; aturannya lihat dampak saat barang habis.",
     "Stok terlalu banyak, stok terlalu sedikit, dan kotak aturan keputusan.",
     ["Sekarang dilemanya. Kalau stok terlalu banyak, biaya penyimpanan dan penanganan naik. Uang perusahaan "
      "tertahan dalam bentuk barang, dan barangnya berisiko rusak atau tidak laku lagi.",
      "Sebaliknya, kalau stok terlalu sedikit, barang bisa habis dan proses operasional terhambat. Di bisnis "
      "ritel online akibatnya lebih cepat terasa. Pelanggan tinggal menutup satu tab dan pindah ke toko lain "
      "dalam hitungan detik.",
      "Jadi berapa yang pas? Aturannya sebenarnya sederhana. Lihat dulu seberapa besar dampaknya kalau barang "
      "sampai habis.",
      "Kalau dampaknya kecil, perusahaan tidak perlu menyimpan banyak stok. Tetapi kalau barangnya lama "
      "tersedia kembali, atau kalau kehabisan barang bisa membuat penjualan hilang, maka stok cadangan perlu "
      "ditambah.",
      "Aturan ini akan kami pakai lagi nanti. Persediaan Amazon itu sangat besar, dan pertanyaan wajarnya "
      "adalah: apakah itu pemborosan, atau justru keputusan yang tepat? Jawabannya akan dibuktikan dengan "
      "aturan tadi."],
     "Untuk menjawabnya, kita perlu tahu dulu kisahnya. Bagian berikutnya akan dibawakan Fitra. Silakan, Fitra.",
     ["Bagian tentang pelanggan yang menutup satu tab membuat contohnya terasa nyata. Boleh diperagakan "
      "dengan gerakan tangan.",
      "Pertanyaan penutup sengaja dibiarkan menggantung. Jangan dijawab di sini, jawabannya ada di slide "
      "Stok Besar Tidak Efisien yang dibawakan pembicara berikutnya.",
      "Ini titik serah-terima. Sebut nama Fitra dengan jelas, lalu mundur satu langkah dan berikan ruang."]),
]

KUNCI = [
    "Slide 1: Persediaan biasanya dianggap beban, tetapi di Amazon justru menjadi keunggulan.",
    "Slide 2: Kami membahas dari strategi dulu, baru teknis.",
    "Slide 3: Nilai Amazon adalah Customer Obsession, semua keputusan berangkat dari pelanggan.",
    "Slide 4: Banyak, cepat, akurat.",
    "Slide 5: Pertanyaan di mana yang paling menentukan; stok didekatkan sebelum pesanan masuk.",
    "Slide 6: Tingkat pelayanan adalah batas, bukan tujuan.",
    "Slide 7: Besar kecilnya stok ditentukan oleh dampak saat barang habis.",
]

QA = [
    ("Kenapa Amazon perlu 150 gudang? Apakah tidak lebih murah punya satu gudang besar?",
     "Satu gudang besar memang lebih murah biaya bangunan dan stoknya, tetapi jarak ke pelanggan menjadi jauh, "
     "sehingga ongkos dan waktu kirim naik. Amazon memilih mendekatkan stok supaya bisa memakai pengiriman yang "
     "murah dan tetap sampai dalam satu sampai dua hari. Rujukan slide 5."),
    ("Apa bedanya objective dan constraint di slide 6?",
     "Objective adalah yang dikejar sampai sekecil mungkin, yaitu total biaya. Constraint adalah batas yang tidak "
     "boleh dilanggar, misalnya janji kirim satu sampai dua hari. Jadi Amazon menekan biaya, tetapi tidak boleh "
     "sampai mengorbankan janji layanannya. Rujukan slide 6."),
    ("Kalau stok besar itu mahal, kenapa Amazon tetap menyimpan banyak?",
     "Karena menurut aturan di slide 7, yang menentukan adalah dampak saat barang habis. Bagi Amazon dampaknya "
     "besar: penjualan hilang dan janji ke pelanggan rusak. Jadi menambah stok masih lebih menguntungkan daripada "
     "menanggung kehabisan barang. Pembuktian lengkapnya ada di bagian Fitra."),
    ("Amazon berdiri 1994 atau 1995?",
     "Perusahaannya didirikan tahun 1994 di Seattle. Situsnya mulai melayani pembeli tahun 1995. Karena itu slide 3 "
     "menyebut 1994, dan slide kisah sukses menyebut rencana bisnis tahun 1995."),
    ("Apa hubungan Customer Obsession dengan pengelolaan persediaan?",
     "Customer Obsession membuat Amazon menetapkan janji layanan lebih dulu: harga murah, kirim cepat, bebas "
     "kesalahan. Janji itulah yang kemudian memaksa Amazon membangun gudang sendiri dan menaruh stok dekat "
     "pelanggan. Jadi persediaan adalah akibat dari nilai perusahaannya, bukan sebaliknya."),
]

ISTILAH = [
    ("Inventory atau persediaan", "barang yang disimpan perusahaan untuk memenuhi permintaan pelanggan."),
    ("Stock out", "kondisi barang habis saat pelanggan ingin membeli."),
    ("Holding cost atau biaya penyimpanan", "biaya menyimpan barang: sewa gudang, uang yang tertahan, dan risiko barang rusak."),
    ("Service level atau tingkat pelayanan", "seberapa sering pesanan bisa dipenuhi langsung dari stok yang ada."),
    ("Lead time atau waktu tunggu", "waktu dari saat memesan sampai barang benar-benar datang."),
    ("Safety stock atau stok cadangan", "stok tambahan sebagai penjaga kalau permintaan naik atau barang telat datang."),
    ("Trade-off", "pilihan yang saling mengorbankan: menambah satu hal berarti mengurangi hal lain."),
    ("Fulfillment center", "gudang Amazon tempat pesanan diambil, dikemas, dan dikirim."),
]

PERUBAHAN = [
    "Tujuan dan Alur Pembahasan naik dari slide 3 menjadi slide 2.",
    "Tentang Amazon turun dari slide 2 menjadi slide 3.",
    "Slide 1, 4, 5, 6, dan 7 tetap pada posisi yang sama.",
    "Akibatnya alurnya berubah: peta pembahasan disampaikan lebih dulu, baru profil perusahaan. "
    "Kalimat transisi pada slide 1, 2, dan 3 sudah disesuaikan dengan urutan ini.",
    "Penyebutan Customer Obsession kini jatuh di slide 3, tetap sebelum Kerangka Berpikir, sehingga benang "
    "merahnya tidak terganggu.",
]


def shade(paragraph, hex_fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def para(doc, text, size=11, bold=False, italic=False, color=BLACK, after=6, before=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size, r.font.color.rgb = Pt(size), color
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


def build():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.left_margin = s.right_margin = Cm(2.2)

    para(doc, "NASKAH PRESENTASI · SLIDE 1–7", 22, bold=True, after=2)
    para(doc, "Inventory Management sebagai Keunggulan Kompetitif Amazon", 14, bold=True, color=MAGENTA, after=2)
    para(doc, "Pembicara: Aulia Sisca Rahmadiyanti  ·  Urutan pembicara: Aulia, Fitra, Bagaskoro", 11, bold=True, after=2)
    para(doc, "Studi Kasus 4 · Operations & Technology Management · Kelompok 4 · Mengikuti urutan slide terbaru (22 slide)",
         10, color=GRAY, after=14)

    para(doc, "Perubahan urutan pada versi ini", 13, bold=True, after=4)
    for t in PERUBAHAN:
        p = doc.add_paragraph(t, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

    para(doc, "Cara memakai naskah ini", 13, bold=True, before=12, after=4)
    for t in ["Naskah ditulis seperti orang berbicara, bukan seperti tulisan laporan. Tidak perlu dihafal kata per "
              "kata; pahami alurnya lalu ucapkan dengan bahasa sendiri.",
              "Setiap slide punya empat bagian: Di layar, Naskah, Transisi ke slide berikutnya, dan Tips.",
              "Kalimat dengan latar merah muda adalah pesan utama slide. Kalau waktu menipis, kalimat itu yang "
              "wajib diucapkan.",
              "Kalimat transisi penting supaya tujuh slide terasa satu alur, bukan tujuh bagian yang terpisah.",
              "Total slide 1 sampai 7 sekitar 8 menit. Tambahkan 1 menit cadangan untuk jeda dan penyesuaian "
              "di ruangan.",
              "Serah-terima ke Fitra dilakukan di akhir slide 7."]:
        p = doc.add_paragraph(t, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

    para(doc, "Ringkasan slide 1–7", 13, bold=True, before=12, after=4)
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Slide", "Judul", "Durasi", "Inti pesan"]):
        c = tbl.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
    for n, title, dur, inti, *_ in SLIDES:
        row = tbl.add_row().cells
        for i, v in enumerate([str(n), title, dur, inti]):
            row[i].text = ""
            r = row[i].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5)
    row = tbl.add_row().cells
    for i, v in enumerate(["", "Total slide 1–7", "± 8 menit", "Seluruhnya dibawakan Aulia; serah-terima ke Fitra di akhir slide 7."]):
        row[i].text = ""
        r = row[i].paragraphs[0].add_run(v)
        r.bold = True
        r.font.size = Pt(9.5)

    para(doc, "Naskah per slide", 16, bold=True, before=20, after=8)
    for n, title, dur, inti, screen, lines, transisi, tips in SLIDES:
        h = para(doc, f"Slide {n} · {title}", 13, bold=True, before=14, after=2)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.keep_together = True
        para(doc, f"Aulia · {dur}", 9.5, bold=True, color=MAGENTA, after=4)
        para(doc, "Di layar: " + screen, 10, italic=True, color=GRAY, after=6)
        for i, line in enumerate(lines):
            if line.startswith("["):
                para(doc, line, 10.5, italic=True, color=GRAY, after=5)
                continue
            key = (i == len(lines) - 1)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line)
            r.font.size = Pt(11.5)
            r.bold = key
            if key:
                shade(p, PINK)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("Transisi ke slide berikutnya: ")
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = p.add_run(transisi)
        r2.italic = True
        r2.font.size = Pt(10.5)
        para(doc, "Tips:", 10, bold=True, color=GRAY, before=2, after=2)
        for t in tips:
            pt = doc.add_paragraph(t, style="List Bullet")
            pt.paragraph_format.space_after = Pt(2)
            for r in pt.runs:
                r.font.size = Pt(10)
                r.italic = True
                r.font.color.rgb = GRAY

    para(doc, "Tujuh kalimat yang cukup dihafal", 16, bold=True, before=20, after=4)
    para(doc, "Kalau gugup dan lupa naskahnya, tujuh kalimat ini sudah cukup untuk membawa slide 1 sampai 7 "
              "secara utuh.", 10.5, color=GRAY, after=8)
    for k in KUNCI:
        p = doc.add_paragraph(k, style="List Number")
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(11.5)

    para(doc, "Persiapan tanya jawab untuk bagian ini", 16, bold=True, before=18, after=4)
    para(doc, "Lima pertanyaan yang paling mungkin muncul dari slide 1 sampai 7, dengan jawaban singkat dan "
              "slide rujukannya. Daftar lengkap 25 pertanyaan ada di dokumen tanya jawab terpisah.",
         10.5, color=GRAY, after=8)
    for i, (q, a) in enumerate(QA, 1):
        para(doc, f"{i}. {q}", 11.5, bold=True, before=6, after=2)
        p = para(doc, a, 11, after=6)
        p.paragraph_format.left_indent = Cm(0.6)

    para(doc, "Istilah dan cara menjelaskannya dengan bahasa sederhana", 13, bold=True, before=14, after=4)
    for term, plain in ISTILAH:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(term + ": ")
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = p.add_run(plain)
        r2.font.size = Pt(10.5)

    doc.save(OUT)
    print("ditulis:", OUT)


if __name__ == "__main__":
    build()
