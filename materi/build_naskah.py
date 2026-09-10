# -*- coding: utf-8 -*-
"""
Naskah presentasi (script presenter) untuk deck final Kelompok 4:
"Inventory Management sebagai Keunggulan Kompetitif Amazon" (21 slide).
Bahasa lisan yang sederhana: mudah dibaca presenter, mudah dipahami audiens.
Jalankan:  python3 build_naskah.py   -> Naskah_Presentasi_Studi_Kasus_4_Amazon_Kelompok_4.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Naskah_Presentasi_Studi_Kasus_4_Amazon_Kelompok_4.docx")
BLACK, GRAY, MAGENTA, PINK = RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x5F, 0x5F, 0x5F), RGBColor(0xC4, 0x18, 0x5C), "F6D5DA"

# ------------------------------------------------------------------ isi naskah
# (nomor slide, judul di layar, pembicara, durasi, "di layar", naskah [paragraf], tips)
SLIDES = [
    (1, "Sampul", "Bagaskoro", "0,5 menit",
     "Judul presentasi, nama anggota Kelompok 4.",
     ["Selamat pagi Ibu Prof. Nur Aini dan teman-teman semua. Kami dari Kelompok 4: Bagaskoro, Aulia Sisca Rahmadiyanti, dan Fitra Aidila.",
      "Hari ini kami membahas studi kasus keempat: Inventory Management sebagai Keunggulan Kompetitif Amazon.",
      "Satu kalimat pembuka dari kami: cara Amazon mengelola stok bukan sekadar urusan gudang. Justru dari situlah Amazon bisa lebih unggul dari pesaingnya. Itu yang akan kami tunjukkan."],
     "Sapa dosen lebih dulu, lalu teman-teman. Bicara pelan di kalimat pertama; audiens masih menyesuaikan diri."),

    (2, "Tujuan dan Alur Pembahasan", "Bagaskoro", "1 menit",
     "Kotak tujuan di atas; enam butir alur pembahasan.",
     ["Ada dua hal yang ingin kami capai. Pertama, memahami mengapa Amazon memilih menyimpan stok sendiri, padahal di awal Amazon tidak mau punya gudang sama sekali. Kedua, mengikuti perjalanan satu pesanan dari saat diklik pelanggan sampai tiba di rumahnya, dan melihat pilihan apa saja yang harus dikorbankan di sepanjang jalan.",
      "Alurnya ada enam bagian. Kerangka berpikir. Kisah sukses Amazon. Proses pemenuhan pesanan. Analisis trade-off. Perkembangan terbaru tahun 2023 dan 2025. Dan terakhir, lesson learned.",
      "Kami sengaja mulai dari strategi dulu, baru masuk ke hal teknis. Alasannya sederhana: keputusan soal stok selalu mengikuti strategi yang sudah dipilih perusahaan, bukan sebaliknya."],
     "Tunjuk keenam butir dengan cepat, tidak perlu dibaca satu per satu."),

    (3, "Inventory Management Amazon (video/gambar)", "Bagaskoro", "1 menit",
     "Video atau gambar suasana di dalam same-day center Amazon.",
     ["Sebelum masuk ke teori, ini gambaran suasana di dalam gudang Amazon.",
      "[Putar video singkat, atau tunjukkan gambarnya beberapa detik.]",
      "Perhatikan tiga hal. Barangnya sangat banyak. Orang dan barang bergerak cepat. Dan semuanya diberi label serta dipindai. Tiga hal itu, banyak, cepat, dan akurat, adalah inti dari seluruh pembahasan kita hari ini."],
     "Kalau video tidak bisa diputar, cukup tunjukkan gambar dan langsung ke kalimat 'Perhatikan tiga hal'. Jangan menunggu lebih dari 30 detik."),

    (4, "Kerangka Berpikir", "Bagaskoro", "1 menit",
     "Tiga pertanyaan: How much, When, Where.",
     ["Setiap perusahaan yang mengelola persediaan harus menjawab tiga pertanyaan dasar. Berapa banyak stok yang disiapkan untuk tiap produk? Kapan harus memesan lagi supaya stok tidak habis? Dan di gudang mana barang disimpan?",
      "Untuk kebanyakan perusahaan, dua pertanyaan pertama yang paling sering dibahas. Untuk Amazon, pertanyaan ketiga, 'di mana', justru yang paling menentukan.",
      "Amazon punya lebih dari 150 gudang. Barang harus sudah ada di gudang yang dekat dengan pelanggan sebelum pelanggan memesan. Kalau barangnya masih jauh, janji kirim satu sampai dua hari tidak mungkin ditepati."],
     "Tekankan kata 'sebelum pelanggan memesan'. Ini kunci yang akan dipakai lagi di slide 10, 13, dan 19."),

    (5, "Struktur Keputusan", "Bagaskoro", "1,5 menit",
     "Objective, Decision Variables, Constraints; kotak tentang tingkat pelayanan.",
     ["Setiap keputusan persediaan punya tiga bagian yang sama.",
      "Bagian pertama, tujuannya. Biasanya menekan tiga biaya: biaya simpan, biaya pesan, dan kerugian saat barang habis. Amazon menambah satu lagi: biaya kalau pesanan salah, karena retur dan tukar barang itu mahal.",
      "Bagian kedua, apa yang harus diputuskan. Berapa yang dipesan, kapan pesan ulang, dan berapa stok cadangan. Amazon menambah satu keputusan lagi: barang apa disimpan di gudang mana.",
      "Bagian ketiga, batasannya. Target layanan, kapasitas gudang, dan waktu tunggu sampai barang datang. Bagi Amazon, batasannya berat: harga harus tetap kompetitif, kirim harus cepat, dan retur harus mudah.",
      "Satu hal penting di slide ini. Tingkat pelayanan bukan tujuan, tetapi batas. Artinya, Amazon tidak berusaha memberi layanan setinggi mungkin dengan biaya berapa pun. Amazon menekan biaya, dengan syarat janji kepada pelanggan tetap dipenuhi."],
     "Kalimat terakhir adalah pesan utama slide ini. Ucapkan pelan, lalu jeda sebentar."),

    (6, "Dilema Persediaan dan Aturan Keputusan", "Bagaskoro", "1 menit",
     "Stok terlalu banyak vs terlalu sedikit; kotak aturan keputusan.",
     ["Sekarang dilemanya. Kalau stok terlalu banyak, biaya simpan naik, uang perusahaan tertahan di gudang, dan barang bisa rusak atau tidak laku lagi. Kalau stok terlalu sedikit, barang habis. Di toko online, pelanggan tinggal klik toko lain dalam hitungan detik.",
      "Jadi, berapa yang pas? Aturannya sederhana. Lihat dulu seberapa besar kerugian kalau barang habis. Kalau kerugiannya kecil, tidak perlu simpan banyak. Kalau kerugiannya besar, misalnya barang lama datangnya atau penjualan bisa hilang, stok cadangan perlu ditambah.",
      "Aturan ini akan kita pakai lagi nanti, untuk menjawab pertanyaan: stok Amazon yang sangat besar itu pemborosan, atau justru keputusan yang benar?"],
     "Pertanyaan penutup dibiarkan menggantung; jawabannya ada di slide 15. Ini penghubung ke pembicara berikutnya."),

    (7, "Transformasi Model Bisnis Amazon", "Aulia", "1 menit",
     "//01 Rencana 1995: pengecer virtual; //02 Kenyataan: gudang terbesar di dunia.",
     ["Sekarang kisahnya. Tahun 1995, rencana Amazon adalah menjadi toko virtual. Tidak punya stok, tidak punya gudang. Pelanggan memesan buku, Amazon meneruskan pesanan itu ke pemasok. Model ini ringan dan murah.",
      "Kenyataannya hari ini sangat berbeda. Amazon punya lebih dari 150 gudang dan 200 juta jenis barang di situsnya. Satu gudang bisa mengeluarkan 200 ribu barang dalam sehari. Dan tenaga kerja yang dipakai untuk satu pesanan kurang dari tiga menit.",
      "Dari 'tanpa gudang' menjadi 'pemilik gudang terbesar di dunia'. Pertanyaannya: apakah rencana awalnya salah?"],
     "Kontraskan kedua sisi dengan tangan: kiri untuk 1995, kanan untuk sekarang."),

    (8, "Apakah Rencana Awalnya Salah?", "Aulia", "1,5 menit",
     "Tiga target Bezos; dua kotak penjelasan.",
     ["Jawabannya: tidak salah, tetapi tidak bisa dipertahankan.",
      "Begitu Amazon berjanji tiga hal kepada pelanggan, yaitu harga terendah, pengiriman tercepat, dan pesanan bebas kesalahan, model tanpa gudang tidak sanggup lagi.",
      "Kenapa? Kalau stok ada di pemasok, kecepatan kirim bergantung pada orang lain. Kalau pemasok salah kirim, yang menanggung biaya retur tetap Amazon, tetapi Amazon tidak bisa mengendalikannya. Dan margin pemasok ikut masuk ke harga.",
      "Jadi, menyimpan stok sendiri bukan berarti Amazon menyimpang dari strategi. Itu konsekuensi dari janji yang sudah dibuat. Begitu janji layanan ditetapkan, gudang sendiri menjadi keharusan."],
     "Tiga target Bezos sebaiknya diucapkan dengan jari: satu, dua, tiga."),

    (9, "Karakter Operasi: 4V", "Aulia", "1,5 menit",
     "Volume, Variety, Variation, Visibility; kotak tentang pengambilan massal.",
     ["Supaya paham kenapa gudang Amazon dirancang seperti itu, kita lihat karakter operasinya lewat 4V.",
      "Volume sangat tinggi: 200 ribu barang per hari per gudang. Dengan volume sebesar ini, otomasi menjadi masuk akal secara biaya.",
      "Variety sangat tinggi: 200 juta jenis barang. Artinya, catatan stok setiap barang harus benar. Salah satu saja, pesanan bisa salah.",
      "Variation tinggi: ada musim puncak, jadi gudang dirancang untuk menangani lonjakan pesanan.",
      "Visibility rendah: pelanggan tidak melihat gudangnya. Mereka hanya melihat harga, waktu tiba, dan isi paket.",
      "Biasanya volume tinggi dan variety tinggi sulit digabung, karena banyak jenis barang membuat proses lambat. Amazon mengatasinya dengan mengambil barang dalam jumlah besar sekaligus, baru dipisah per pesanan di ujung proses."],
     "Kalau waktu mepet, cukup sebut Volume dan Variety, lalu langsung ke kalimat terakhir."),

    (10, "Proses Pemenuhan Pesanan: Langkah 1 sampai 4", "Aulia", "1,5 menit",
     "//01 Pesanan masuk; //02 Flow meister; //03 Picking; //04 Konveyor.",
     ["Sekarang kita ikuti perjalanan satu pesanan.",
      "Langkah satu, pesanan masuk. Sistem menentukan gudang mana yang punya barangnya. Ini jawaban dari pertanyaan 'di mana' tadi. Dan ini hanya bisa jalan kalau catatan stok setiap gudang akurat.",
      "Langkah dua, pesanan diterima oleh flow meister, semacam pengatur lalu lintas di gudang. Ia memecah pesanan menjadi tugas-tugas pengambilan barang dan menentukan siapa mengambil apa.",
      "Langkah tiga, picking atau pengambilan barang. Pekerja mengambil banyak barang untuk beberapa pesanan sekaligus dalam satu perjalanan. Cepat, tetapi ada risiko barang tertukar. Amazon mengurangi risiko itu dengan robot Kiva, yang membawa rak langsung ke pekerja. Nanti kita bahas lebih detail.",
      "Langkah empat, barang dimasukkan ke kotak dan bergerak di konveyor sepanjang lebih dari 10 mil. Di sepanjang jalan, bar code dipindai sampai 15 kali supaya barang selalu terlacak."],
     "Ikuti urutan nomor di layar dengan tangan, dari //01 ke //04."),

    (11, "Proses Pemenuhan Pesanan: Langkah 5 sampai 8", "Aulia", "1,5 menit",
     "//05 Dikelompokkan; //06 Bungkus kado; //07 Dikemas, ditimbang, dilabeli; //08 Tiba di pelanggan.",
     ["Langkah lima, barang dikelompokkan kembali sesuai pesanan. Bar code dicocokkan dengan nomor pesanan, sehingga semua isi satu pesanan berkumpul di satu kotak.",
      "Langkah enam, pembungkusan kado. Ini masih dikerjakan tangan, sekitar 30 paket per jam, karena bentuk barang beragam dan jumlahnya kecil. Pelajarannya: tidak semua hal perlu diotomasi.",
      "Langkah tujuh, paket ditimbang. Kalau beratnya tidak sesuai, berarti isinya ada yang salah. Ini pemeriksaan terakhir. Setelah itu paket diberi label dan diserahkan ke pengirim. Sekitar 60 persen paket dikirim lewat USPS, layanan pos Amerika yang biayanya murah.",
      "Langkah delapan, paket tiba dalam satu atau dua hari. Ingat, pelanggan hanya menilai tiga hal: cepat, tepat waktu, dan isinya benar."],
     "Setelah slide ini, serahkan ke pembicara berikutnya dengan kalimat: 'Dari proses ini muncul dua angka yang sering disebut. Fitra akan menjelaskannya.'"),

    (12, "Dua Angka dengan Arti yang Berbeda", "Fitra", "1 menit",
     "Angka besar: < 3 menit dan 1–2 hari; dua kotak penjelasan.",
     ["Ada dua angka yang sering disebut tentang Amazon, dan sering tertukar. Kurang dari tiga menit, dan satu sampai dua hari.",
      "Tiga menit adalah waktu kerja orang untuk satu pesanan di dalam gudang. Ini ukuran biaya. Semakin singkat, semakin murah.",
      "Satu sampai dua hari adalah waktu tunggu pelanggan, dari pesan sampai barang diterima. Ini ukuran kecepatan.",
      "Yang pertama hasil dari proses gudang yang efisien. Yang kedua hasil dari stok yang ditaruh dekat pelanggan. Jadi keduanya jangan dicampur. Amazon unggul di dua-duanya, tetapi mekanismenya berbeda."],
     "Tunjuk angka kiri saat menyebut 'biaya', angka kanan saat menyebut 'kecepatan'."),

    (13, "Mekanisme Robot Kiva", "Fitra", "1,5 menit",
     "Empat butir angka; Dampak pada persediaan.",
     ["Dari mana angka tiga menit itu? Dari robot Kiva.",
      "Lebih dari 10 ribu robot membawa rak ke pekerja. Jadi bukan orang yang berjalan ke barang, tetapi barang yang datang ke orang. Ini menghemat sekitar 20 mil jalan kaki per pekerja per hari.",
      "Hasilnya, kecepatan ambil dan pindai naik dari 100 menjadi 300 barang per jam. Biaya per pesanan turun 20 sampai 40 persen, atau kira-kira 400 sampai 900 juta dolar per tahun.",
      "Dampaknya ke persediaan juga besar. Barang tidak harus punya tempat rak yang tetap, dan setiap kali diambil langsung dipindai. Catatan stok jadi selalu terbaru.",
      "Inilah yang membuat langkah satu tadi, memilih gudang mana yang menangani pesanan, bisa dipercaya. Data stoknya akurat."],
     "Gerakan tangan: 'barang datang ke orang' bisa diperagakan dengan menarik tangan ke arah dada."),

    (14, "Mengukur Produktivitas secara Utuh", "Fitra", "1 menit",
     "Single-factor +200%; Multi-factor biaya per pesanan turun 20–40%; dampak tahunan.",
     ["Sekarang hati-hati dengan angka 200 persen.",
      "Dari 100 ke 300 barang per jam memang naik 200 persen. Tetapi itu hanya melihat hasil kerja orang, belum menghitung biaya beli dan biaya operasi robotnya. Ini yang disebut produktivitas satu faktor.",
      "Kalau biaya robot ikut dihitung, penghematannya sekitar 0,70 sampai 1,50 dolar per pesanan. Kelihatannya kecil. Tetapi dikalikan ratusan juta pesanan, jadi 400 sampai 900 juta dolar setahun.",
      "Jadi kalau ingin jujur menilai produktivitas, angka yang dipakai adalah penurunan biaya per pesanan, bukan angka 200 persen."],
     "Ini titik yang biasanya ditanyakan dosen. Pastikan bisa membedakan 'satu faktor' dan 'banyak faktor' dengan kalimat sendiri."),

    (15, "Stok Besar, Tidak Efisien?", "Fitra", "1,5 menit",
     "Persediaan sampai 50% modal; kotak tujuan operasi; Yang diminimalkan Amazon.",
     ["Sekarang pertanyaan yang tadi kita gantung: stok Amazon besar sekali, bukankah itu tidak efisien? Nilai persediaan bisa sampai 50 persen dari modal, tersebar di lebih dari 150 gudang.",
      "Jawabannya kembali ke tujuan. Amazon tidak berusaha punya stok sesedikit mungkin. Amazon berusaha menekan total biaya sambil tetap memenuhi janji ke pelanggan.",
      "Kalau barang habis, Amazon bukan hanya kehilangan penjualan, tetapi juga mengecewakan pelanggan. Menurut aturan di slide 6 tadi, kerugian barang habis di sini besar, jadi stok besar itu justru sejalan dengan strateginya.",
      "Yang ditekan Amazon adalah dua hal lain: waktu kerja per pesanan, kurang dari tiga menit, dan kesalahan, mendekati nol.",
      "Kesimpulannya: bukan stok sedikit, tetapi stok besar yang dikelola cepat dan akurat."],
     "Hubungkan kembali ke slide 6 secara eksplisit; audiens akan merasa alurnya utuh."),

    (16, "Lima Tujuan Kinerja Operasi", "Fitra", "1,5 menit",
     "Speed, Quality, Dependability, Flexibility, Cost; kotak penutup.",
     ["Kalau dipetakan ke lima tujuan kinerja operasi, semuanya terpenuhi.",
      "Speed: pesanan tiba satu sampai dua hari, dan proses di gudang kurang dari tiga menit.",
      "Quality: kesalahan hampir nol, berkat 15 kali pemindaian dan penimbangan sebelum kirim.",
      "Dependability: janji satu sampai dua hari benar-benar ditepati. Bukan sekadar cepat sesekali, tetapi cepat seperti yang dijanjikan.",
      "Flexibility: 70 persen pesanan berisi lebih dari satu barang, dan sistemnya sanggup menangani berbagai kombinasi.",
      "Cost: harga tetap kompetitif, 60 persen paket lewat USPS yang murah, dan biaya proses turun 20 sampai 40 persen.",
      "Yang menarik, kelimanya didukung oleh sistem yang sama: data stok akurat, proses singkat, dan pemeriksaan berlapis."],
     "Sebut kelima kata bahasa Inggrisnya persis seperti di layar, karena itu istilah baku di mata kuliah ini."),

    (17, "Tiga Trade-off Utama", "Bagaskoro", "1,5 menit",
     "Tiga trade-off dengan aturan keputusannya.",
     ["Tidak ada keputusan yang gratis. Setiap pilihan punya konsekuensi di sisi lain. Ada tiga trade-off utama di kasus ini.",
      "Pertama, biaya simpan lawan tingkat layanan. Biaya simpan sekitar 26 persen per tahun dari nilai barang. Aturannya: tambah stok kalau kerugian barang habis lebih besar daripada biaya simpannya. Untuk toko online, syarat ini terpenuhi.",
      "Kedua, otomasi lawan kesalahan dan tenaga kerja. Aturannya: otomasi layak kalau penghematan per pesanan dikali jumlah pesanan setahun lebih besar daripada biaya otomasinya. Penentunya adalah volume penjualan.",
      "Ketiga, kecepatan kirim lawan biaya kirim. Amazon tidak membayar kurir mahal supaya cepat. Kecepatan dibangun dari dalam: stok didekatkan ke pelanggan dan proses gudang dibuat cepat, sehingga pos yang murah pun sampai dalam satu dua hari."],
     "Setiap trade-off punya 'aturan'. Ucapkan kata 'aturannya' dengan jelas; itu cara berpikir yang dinilai."),

    (18, "Verifikasi Trade-off dan Efficient Frontier", "Bagaskoro", "1,5 menit",
     "Empat pasangan tujuan; dua butir efficient frontier.",
     ["Sebelum menyebut sesuatu sebagai trade-off, kita perlu cek dulu: dua tujuan itu benar bertentangan, atau sebenarnya sejalan?",
      "Akurasi dan biaya: sejalan. Pesanan yang benar mengurangi biaya retur dan tukar barang.",
      "Kecepatan proses dan biaya: sejalan. Proses yang singkat berarti biaya kerja kecil.",
      "Stok dan biaya simpan: ini trade-off sungguhan. Amazon memilihnya dengan sadar.",
      "Kecepatan kirim dan biaya kirim: trade-off juga, tetapi Amazon menghindarinya lewat penempatan stok dan proses gudang yang cepat.",
      "Kalau digambarkan sebagai efficient frontier, yaitu batas terbaik antara biaya dan layanan: model tanpa gudang memang murah, tetapi tidak sanggup memenuhi target layanan Amazon. Gudang manual kalah pada volume sebesar Amazon. Gudang otomatis berada di batas terbaik untuk skala Amazon.",
      "Satu catatan: untuk perusahaan dengan volume kecil, gudang manual bisa saja yang lebih efisien. Jadi jawabannya bergantung pada skala."],
     "Jelaskan 'efficient frontier' dengan kata sederhana dulu ('batas terbaik'), baru sebut istilahnya."),

    (19, "Perkembangan Terbaru", "Aulia", "1 menit",
     "2023 Regionalisasi jaringan; 2025 Skala robot.",
     ["Dua perkembangan terbaru.",
      "Tahun 2023, Amazon membagi wilayah Amerika menjadi delapan region, masing-masing dengan stok lengkap. 76 persen permintaan dipenuhi dari dalam region. Fasilitas same-day menyimpan 100 ribu barang terlaris. Ini pertanyaan 'di mana' tadi, sekarang dijawab per wilayah: barang yang paling laris ditaruh paling dekat.",
      "Tahun 2025, robotnya sudah lebih dari satu juta. 75 persen pengiriman melibatkan robot, dan paket per karyawan naik dari 175 menjadi 3.870 per tahun.",
      "Tetapi ingat pelajaran dari slide 14. Angka per karyawan itu produktivitas satu faktor. Belum menghitung biaya investasi dan operasi robotnya."],
     "Kalimat terakhir menunjukkan kelompok konsisten dengan cara berpikirnya sendiri. Jangan dilewati."),

    (20, "Lesson Learned", "Fitra", "1,5 menit",
     "Enam pelajaran //01 sampai //06.",
     ["Enam pelajaran yang bisa dibawa pulang.",
      "Satu, persediaan bisa menjadi keunggulan bersaing, asalkan jumlah dan lokasinya dikelola dengan baik.",
      "Dua, data stok harus akurat dulu. Sistem canggih tidak ada gunanya kalau datanya salah.",
      "Tiga, kecepatan dibangun dari dalam: dari posisi stok yang dekat dan proses yang singkat.",
      "Empat, pengecekan berlapis tetap layak, selama biaya kesalahan lebih besar daripada biaya pengecekan.",
      "Lima, otomasi butuh volume besar. Penghematan kecil per pesanan baru terasa kalau pesanannya banyak.",
      "Enam, trade-off yang tersisa harus dipilih dengan sadar. Amazon menerima biaya simpan yang besar supaya barang selalu tersedia."],
     "Baca keenamnya dengan ritme tetap. Ini bagian yang paling mungkin diingat audiens."),

    (21, "Terima Kasih", "Fitra", "0,5 menit",
     "Slide penutup.",
     ["Itu presentasi dari kami. Kalau harus diringkas dalam satu kalimat: Amazon unggul bukan karena stoknya sedikit, tetapi karena stoknya besar, akurat, dan dekat dengan pelanggan.",
      "Terima kasih, Bu. Kami terbuka untuk pertanyaan dan diskusi."],
     "Semua anggota berdiri saat sesi tanya jawab. Sepakati dulu siapa menjawab topik apa (lihat halaman terakhir)."),
]

QA = [
    ("Kenapa Amazon tidak menitipkan saja stoknya ke pihak ketiga supaya lebih murah?",
     "Karena tiga janji Amazon, yaitu harga terendah, kirim tercepat, dan bebas kesalahan, tidak bisa dijamin kalau stok dan prosesnya ada di tangan orang lain. Kesalahan pihak ketiga tetap menjadi biaya retur Amazon, tetapi tidak bisa dikendalikan. Slide 8."),
    ("Bukankah stok sebesar itu pemborosan?",
     "Tergantung tujuannya. Amazon tidak meminimalkan stok, tetapi meminimalkan total biaya pada tingkat layanan yang dijanjikan. Karena kerugian saat barang habis besar, stok besar justru sejalan dengan strategi. Yang ditekan adalah waktu kerja per pesanan dan kesalahan. Slide 6 dan 15."),
    ("Apa bedanya angka 200 persen dan angka 20 sampai 40 persen?",
     "200 persen adalah produktivitas satu faktor, hanya hasil kerja orang. 20 sampai 40 persen adalah penurunan biaya per pesanan setelah biaya robot dihitung. Angka kedua yang lebih jujur. Slide 14."),
    ("Apakah cara Amazon bisa ditiru perusahaan kecil?",
     "Prinsipnya bisa: data stok akurat dulu, stok didekatkan ke pelanggan, dan pengecekan dilakukan selama biaya kesalahan lebih besar. Otomasinya belum tentu, karena otomasi baru layak pada volume besar. Pada volume kecil, gudang manual bisa lebih efisien. Slide 17, 18, dan 20."),
    ("Apa hubungan kasus ini dengan teori persediaan di buku?",
     "Fungsi persediaan untuk permintaan yang diantisipasi, akurasi catatan lewat sistem perpetual dengan bar code, analisis ABC pada 100 ribu barang terlaris, dan keputusan penempatan barang di gudang. Semuanya muncul di kasus ini, dengan urutan: akurasi dulu, baru model."),
]

# ------------------------------------------------------------------ pembuat dokumen
def shade(paragraph, hex_fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def para(doc, text, size=11, bold=False, italic=False, color=BLACK, after=6, before=0, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold, r.italic = bold, italic
    r.font.size, r.font.color.rgb = Pt(size), color
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
    if align: p.alignment = align
    return p


def build():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.0); s.left_margin = s.right_margin = Cm(2.2)

    # ---- judul
    para(doc, "NASKAH PRESENTASI", 22, bold=True, after=2)
    para(doc, "Inventory Management sebagai Keunggulan Kompetitif Amazon", 14, bold=True, color=MAGENTA, after=2)
    para(doc, "Studi Kasus 4 · Operations & Technology Management · Kelompok 4: Bagaskoro, Aulia Sisca Rahmadiyanti, Fitra Aidila", 10, color=GRAY, after=12)

    para(doc, "Cara memakai naskah ini", 13, bold=True, before=6, after=4)
    for t in [
        "Naskah ditulis dalam bahasa lisan. Tidak perlu dihafal kata per kata; pahami alurnya, lalu ucapkan dengan kalimat sendiri.",
        "Tiap slide punya tiga bagian: 'Di layar' (apa yang sedang dilihat audiens), 'Naskah' (yang diucapkan), dan 'Tips' (cara membawakan).",
        "Kalimat yang dicetak tebal adalah pesan utama slide. Kalau waktu habis, kalimat tebal itu yang wajib diucapkan.",
        "Total durasi sekitar 26 menit. Kalau jatah waktu 15 menit, pakai kolom 'Versi singkat' di tabel pembagian.",
        "Pembagian pembicara di bawah hanya usulan; silakan ditukar. Pastikan setiap orang tahu slide serah-terimanya.",
    ]:
        p = doc.add_paragraph(t, style="List Bullet"); p.paragraph_format.space_after = Pt(3)

    # ---- tabel pembagian
    para(doc, "Pembagian pembicara dan waktu (usulan)", 13, bold=True, before=12, after=4)
    tbl = doc.add_table(rows=1, cols=5); tbl.style = "Light Grid Accent 1"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["Slide", "Judul", "Pembicara", "Durasi", "Versi singkat (15 menit)"]
    for i, h in enumerate(hdr):
        c = tbl.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
    short = {1: "0,5", 2: "0,5", 3: "lewati", 4: "1", 5: "1", 6: "0,5", 7: "0,5", 8: "1", 9: "0,5", 10: "1", 11: "1", 12: "0,5",
             13: "1", 14: "1", 15: "1", 16: "0,5", 17: "1", 18: "1", 19: "0,5", 20: "1", 21: "0,5"}
    for n, title, who, dur, *_ in SLIDES:
        row = tbl.add_row().cells
        for i, v in enumerate([str(n), title, who, dur, short[n] + (" menit" if short[n] != "lewati" else "")]):
            row[i].text = ""; r = row[i].paragraphs[0].add_run(v); r.font.size = Pt(9.5)
    row = tbl.add_row().cells
    for i, v in enumerate(["", "Total", "", "± 26 menit", "± 15 menit"]):
        row[i].text = ""; r = row[i].paragraphs[0].add_run(v); r.bold = True; r.font.size = Pt(9.5)
    para(doc, "Serah-terima: slide 6 ke 7 (Bagaskoro ke Aulia), slide 11 ke 12 (Aulia ke Fitra), slide 16 ke 17 (Fitra ke Bagaskoro), slide 18 ke 19 (Bagaskoro ke Aulia), slide 19 ke 20 (Aulia ke Fitra).", 9.5, color=GRAY, before=6)

    # ---- naskah per slide
    para(doc, "Naskah per slide", 16, bold=True, before=18, after=8)
    for n, title, who, dur, screen, lines, tips in SLIDES:
        h = para(doc, f"Slide {n} · {title}", 13, bold=True, before=14, after=2)
        h.paragraph_format.keep_with_next = True
        para(doc, f"{who} · {dur}", 9.5, color=MAGENTA, bold=True, after=4)
        p = para(doc, "Di layar: " + screen, 10, italic=True, color=GRAY, after=6)
        for i, line in enumerate(lines):
            key = (i == len(lines) - 1)   # kalimat terakhir = pesan utama slide
            if line.startswith("["):
                para(doc, line, 10.5, italic=True, color=GRAY, after=5)
                continue
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line); r.font.size = Pt(11.5); r.bold = key
            if key: shade(p, PINK)
        t = para(doc, "Tips: " + tips, 10, italic=True, color=GRAY, before=2, after=4)

    # ---- tanya jawab
    doc.add_page_break()
    para(doc, "Persiapan tanya jawab", 16, bold=True, after=4)
    para(doc, "Lima pertanyaan yang paling mungkin muncul, dengan jawaban singkat dan slide rujukannya. Usulan pembagian: pertanyaan tentang strategi dan trade-off dijawab Bagaskoro, tentang proses dan perkembangan terbaru oleh Aulia, tentang angka dan produktivitas oleh Fitra.", 10.5, color=GRAY, after=8)
    for i, (q, a) in enumerate(QA, 1):
        para(doc, f"{i}. {q}", 11.5, bold=True, before=6, after=2)
        p = para(doc, a, 11, after=6); p.paragraph_format.left_indent = Cm(0.6)

    para(doc, "Istilah yang sebaiknya dijelaskan dengan kata sederhana", 13, bold=True, before=14, after=4)
    for term, plain in [
        ("Holding cost / carrying cost", "biaya menyimpan barang: sewa gudang, modal tertahan, risiko rusak. Sekitar 26 persen per tahun dari nilai barang."),
        ("Service level", "tingkat layanan: seberapa sering pesanan bisa dipenuhi dari stok yang ada."),
        ("Trade-off", "pilihan yang saling mengorbankan: menambah yang satu berarti mengurangi yang lain."),
        ("Single-factor vs multi-factor productivity", "produktivitas satu faktor hanya melihat satu input, misalnya jam kerja; multi-factor menghitung semua input, termasuk biaya robot."),
        ("Efficient frontier", "batas terbaik: kombinasi biaya dan layanan yang tidak bisa diperbaiki lagi tanpa mengorbankan salah satunya."),
        ("Flow meister", "pengatur alur di gudang, yang membagi pesanan menjadi tugas-tugas pengambilan barang."),
        ("Picking / batch picking", "mengambil barang dari rak; batch picking berarti mengambil untuk beberapa pesanan sekaligus dalam satu perjalanan."),
    ]:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(term + ": "); r.bold = True; p.add_run(plain)

    doc.save(OUT); print("ditulis:", OUT)


if __name__ == "__main__":
    build()
