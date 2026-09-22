"""Design tokens and slide-building helpers for the Strategic Management deck."""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- tokens ----
INK        = RGBColor(0x1A, 0x16, 0x14)   # near-black, warm
INK_SOFT   = RGBColor(0x4A, 0x43, 0x3F)   # secondary text
MUTED      = RGBColor(0x8C, 0x83, 0x7E)   # tertiary text / de-emphasis mark
RED        = RGBColor(0xDA, 0x29, 0x1C)   # Manchester United red
RED_DEEP   = RGBColor(0x8F, 0x1A, 0x11)   # darker red, for depth
GOLD       = RGBColor(0xB0, 0x84, 0x32)   # accent for conclusions
CREAM      = RGBColor(0xFB, 0xF9, 0xF7)   # slide surface
PAPER      = RGBColor(0xFF, 0xFF, 0xFF)   # card surface
RULE       = RGBColor(0xE4, 0xDF, 0xDA)   # hairlines
TINT       = RGBColor(0xF4, 0xEF, 0xEA)   # subtle fill
RED_TINT   = RGBColor(0xFA, 0xEC, 0xEA)   # subtle red fill
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

H_FONT = "Calibri"      # headings + body: ships with Office on Windows and Mac
S_FONT = "Georgia"      # serif, for section numerals and pull-quotes

SW, SH = 13.333, 7.5    # slide size in inches
ML, MR = 0.72, 0.72     # side margins
CW = SW - ML - MR       # content width = 11.893"


# ------------------------------------------------------------- primitives ---
def set_indent(p, mar_l=0.0, first_line=0.0):
    """Set paragraph indent in inches (python-pptx exposes no paragraph_format)."""
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(mar_l * 914400)))
    pPr.set('indent', str(int(first_line * 914400)))


def rect(slide, l, t, w, h, fill=None, line=None, line_w=0.75):
    """Plain rectangle. Dimensions in inches."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t),
                                 Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.text_frame.word_wrap = True
    return shp


def txt(slide, text, l, t, w, h, size=14, color=INK, bold=False, font=H_FONT,
        align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP, italic=False,
        space_after=0, caps=False):
    """Single-paragraph text box. Dimensions in inches, size in points."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    if caps:
        r.font._rPr.set('cap', 'all')
    return box


def para(box, text, size=14, color=INK, bold=False, font=H_FONT,
         align=PP_ALIGN.LEFT, spacing=1.0, space_before=0, space_after=0,
         italic=False, indent=0.0):
    """Append a paragraph to an existing text box."""
    tf = box.text_frame
    p = tf.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if indent:
        set_indent(p, indent, 0.0)
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    return p


def rich(box, parts, size=14, color=INK, font=H_FONT, align=PP_ALIGN.LEFT,
         spacing=1.0, space_before=0, space_after=0, first=False):
    """Append a paragraph made of (text, bold, color) runs.

    parts: list of str, or (str, bold), or (str, bold, RGBColor).
    """
    tf = box.text_frame
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    for part in parts:
        if isinstance(part, str):
            t, b, c = part, False, color
        elif len(part) == 2:
            t, b, c = part[0], part[1], color
        else:
            t, b, c = part
        r = p.add_run()
        r.text = t
        f = r.font
        f.name, f.size, f.bold = font, Pt(size), b
        f.color.rgb = c
    return p


def bullets(slide, items, l, t, w, h, size=14, color=INK, gap=7, spacing=1.16,
            marker="—", marker_color=RED, bold_lead=True):
    """Bulleted block. Each item: str, or (lead, rest) where lead renders bold."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        set_indent(p, 0.26, -0.26)
        rm = p.add_run()
        rm.text = marker + "  "
        rm.font.name, rm.font.size, rm.font.bold = H_FONT, Pt(size), True
        rm.font.color.rgb = marker_color
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run()
            r1.text = lead
            r1.font.name, r1.font.size, r1.font.bold = H_FONT, Pt(size), bold_lead
            r1.font.color.rgb = color
            r2 = p.add_run()
            r2.text = rest
            r2.font.name, r2.font.size = H_FONT, Pt(size)
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = item
            r.font.name, r.font.size = H_FONT, Pt(size)
            r.font.color.rgb = color
    return box


# ------------------------------------------------------------ table borders --
def set_cell_border(cell, edges=("L", "R", "T", "B"), color=RULE, width=0.75):
    """Draw borders on a table cell. edges from L/R/T/B.

    CT_TableCellProperties requires lnL/lnR/lnT/lnB to precede the fill
    element, so the line elements are always (re)inserted at the front of
    tcPr in that exact order -- appending them breaks the schema and the
    file is then rejected by strict readers.
    """
    from lxml import etree
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    wanted = {tag_map_key: None for tag_map_key in order}
    tag_of = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}
    tcPr = cell._tc.get_or_add_tcPr()

    # drop any existing line elements, keeping the ones we are not redrawing
    existing = {}
    for tag in order:
        el = tcPr.find(qn(tag))
        if el is not None:
            tcPr.remove(el)
            existing[tag] = el

    for e in edges:
        tag = tag_of[e]
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(width * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
        tcPr.remove(ln)
        existing[tag] = ln

    for i, tag in enumerate(order):
        el = existing.get(tag)
        if el is not None:
            tcPr.insert(i, el)


def style_cell(cell, text, size=12, color=INK, bold=False, fill=None,
               align=PP_ALIGN.LEFT, font=H_FONT, anchor=MSO_ANCHOR.MIDDLE,
               spacing=1.06):
    cell.margin_left = Inches(0.10)
    cell.margin_right = Inches(0.10)
    cell.margin_top = Inches(0.055)
    cell.margin_bottom = Inches(0.055)
    cell.vertical_anchor = anchor
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = font, Pt(size), bold
    r.font.color.rgb = color


def make_table(slide, rows, l, t, w, col_w, row_h=0.36, head_h=0.40,
               head_fill=INK, head_color=WHITE, size=12, head_size=11,
               aligns=None, zebra=True, body_fills=None):
    """Build a table. rows[0] is the header. col_w: list of inches."""
    n_r, n_c = len(rows), len(rows[0])
    heights = [head_h] + [row_h] * (n_r - 1)
    tbl_shape = slide.shapes.add_table(n_r, n_c, Inches(l), Inches(t),
                                       Inches(w), Inches(sum(heights)))
    tbl = tbl_shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for i, hh in enumerate(heights):
        tbl.rows[i].height = Inches(hh)
    aligns = aligns or [PP_ALIGN.LEFT] * n_c
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            if ri == 0:
                style_cell(cell, val, size=head_size, color=head_color,
                           bold=True, fill=head_fill, align=aligns[ci])
                set_cell_border(cell, ("T", "B", "L", "R"), head_fill, 0.75)
            else:
                if body_fills is not None:
                    f = body_fills[ri - 1]
                elif zebra:
                    f = PAPER if ri % 2 else TINT
                else:
                    f = PAPER
                style_cell(cell, val, size=size, color=INK,
                           bold=(ci == 0), fill=f, align=aligns[ci])
                set_cell_border(cell, ("T", "B", "L", "R"), RULE, 0.75)
    return tbl


# --------------------------------------------------------------- chrome -----
class Deck:
    def __init__(self, prs, total):
        self.prs = prs
        self.total = total
        self.n = 0

    def blank(self, bg=CREAM):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        if bg is not None:
            bgr = rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1, fill=bg)
            bgr.element.getparent().remove(bgr.element)
            s.shapes._spTree.insert(2, bgr.element)
        return s

    def footer(self, s, label, color=MUTED, rule=RULE):
        rect(s, ML, SH - 0.62, CW, 0.012, fill=rule)
        txt(s, label, ML, SH - 0.47, CW - 1.2, 0.28, size=9.5, color=color,
            spacing=1.0, caps=True)
        txt(s, "%02d" % self.n, SW - MR - 1.0, SH - 0.47, 1.0, 0.28,
            size=9.5, color=color, align=PP_ALIGN.RIGHT, bold=True)

    def head(self, s, kicker, title, sub=None, sublines=1, kcolor=RED, tsize=30):
        """Standard content-slide header. Returns the y where content starts.

        sublines must be the real wrapped line count of `sub` (roughly
        one line per 125 characters at 13.5pt across the content width).
        """
        rect(s, ML, 0.52, 0.30, 0.055, fill=kcolor)
        txt(s, kicker, ML + 0.42, 0.44, CW - 0.42, 0.26, size=10, color=kcolor,
            bold=True, spacing=1.0, caps=True)
        txt(s, title, ML, 0.78, CW, 0.62, size=tsize, color=INK, bold=True,
            spacing=0.98)
        y = 1.50
        if sub:
            txt(s, sub, ML, y, CW - 0.40, 0.30 * sublines, size=13.5,
                color=INK_SOFT, spacing=1.18)
            y += 0.30 * sublines + 0.26
        return y


# ------------------------------------------------------------ layout band ---
# Content must end by CBOT_NOTE when a bottom note follows, or by CBOT when
# it does not. The footer rule sits at SH - 0.62 = 6.88.
CBOT_NOTE = 5.84
CBOT = 6.72
NOTE_TOP = 5.94


def note(s, parts, dark=False, size=12, top=NOTE_TOP, h=0.76,
         left=ML, width=CW):
    """Bottom callout band. Holds two lines of `size`-pt text at h=0.76."""
    rect(s, left, top, width, h, fill=INK if dark else RED_TINT)
    rect(s, left, top, 0.05, h, fill=RED)
    b = txt(s, "", left + 0.30, top + 0.15, width - 0.62, h - 0.26, size=size)
    rich(b, parts, size=size, first=True, spacing=1.16,
         color=WHITE if dark else INK)
    return b
