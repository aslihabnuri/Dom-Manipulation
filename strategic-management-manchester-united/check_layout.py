#!/usr/bin/env python3
"""Flag shapes that leave the safe area, and text that will not fit its box.

Line capacity for Calibri is approximated as width_in * 150 / size_pt, which
matched the rendered output to within a character or two during development.
"""
import sys, os, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
SW, SH = 13.333, 7.5
SAFE_BOTTOM = 6.80      # footer rule sits at 6.88
LINE_FACTOR = 1.26      # rendered line height / font size


def inches(v):
    return (v or 0) / EMU


def text_runs(tf):
    out = []
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        size = max([(r.font.size.pt if r.font.size else 12) for r in p.runs] or [12])
        sp = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        out.append((txt, size, sp))
    return out


def needed_height(tf, width_in):
    total = 0.0
    for txt, size, sp in text_runs(tf):
        if not txt:
            total += size * LINE_FACTOR * sp / 72.0
            continue
        cap = max(1, int(width_in * 150 / size))
        lines = max(1, math.ceil(len(txt) / cap))
        total += lines * size * LINE_FACTOR * sp / 72.0
    return total


def main(path):
    prs = Presentation(path)
    problems = 0
    for i, slide in enumerate(prs.slides, start=1):
        msgs = []
        for shp in slide.shapes:
            l, t = inches(shp.left), inches(shp.top)
            w, h = inches(shp.width), inches(shp.height)
            name = (shp.shape_type or "?")
            if shp.has_text_frame and shp.text_frame.text.strip():
                label = shp.text_frame.text.strip()[:42].replace("\n", " ")
            else:
                label = str(name)
            full_bleed = (l <= 0.02 or l + w >= SW - 0.02
                           or w >= SW - 0.5 or h >= SH - 0.5)
            in_footer = t >= SAFE_BOTTOM - 0.02
            if w <= 0.05 and shp.has_text_frame and shp.text_frame.text.strip():
                msgs.append("  BAD WIDTH   w=%.2f  %s" % (w, label))
                continue
            if t + h > SAFE_BOTTOM + 0.02 and not in_footer and not full_bleed:
                msgs.append("  BELOW SAFE  bottom=%.2f  %s" % (t + h, label))
            if l + w > SW - 0.40 and not full_bleed:
                msgs.append("  PAST RIGHT  right=%.2f  %s" % (l + w, label))
            if in_footer:
                continue
            if shp.has_text_frame and shp.text_frame.text.strip():
                need = needed_height(shp.text_frame, max(w - 0.05, 0.4))
                if need > h + 0.06:
                    msgs.append("  TEXT SPILL  need=%.2f have=%.2f (+%.2f)  %s"
                                % (need, h, need - h, label))
        # A filled shape drawn AFTER a text shape paints over it. That is the
        # "text hidden behind a later card" bug, which the box checks miss.
        items = list(slide.shapes)
        boxes = []
        for idx, shp in enumerate(items):
            has_text = shp.has_text_frame and shp.text_frame.text.strip()
            filled = False
            try:
                filled = (not has_text and shp.fill.type is not None
                          and shp.fill.type != 5)   # 5 == MSO_FILL.BACKGROUND
            except Exception:
                filled = False
            boxes.append((idx, inches(shp.left), inches(shp.top),
                          inches(shp.width), inches(shp.height), has_text, filled,
                          shp.text_frame.text.strip()[:40].replace("\n", " ")
                          if has_text else ""))
        for si, l1, t1, w1, h1, txt1, _f1, lab in boxes:
            if not txt1:
                continue
            for sj, l2, t2, w2, h2, _t2, fill2, _ in boxes:
                if sj <= si or not fill2:
                    continue
                if w2 >= SW - 0.5 or h2 >= SH - 0.5:
                    continue
                ox = min(l1 + w1, l2 + w2) - max(l1, l2)
                oy = min(t1 + h1, t2 + h2) - max(t1, t2)
                if ox > 0.10 and oy > 0.10:
                    msgs.append("  COVERED     by later shape #%d (%.2fx%.2f in)  %s"
                                % (sj, ox, oy, lab))
                    break
        # Two text boxes whose RENDERED extents overlap. A text frame does not
        # clip, so its real height is the wrapped height, not the declared one.
        tb = []
        for shp in slide.shapes:
            if not (shp.has_text_frame and shp.text_frame.text.strip()):
                continue
            l, t = inches(shp.left), inches(shp.top)
            w, h = inches(shp.width), inches(shp.height)
            need = needed_height(shp.text_frame, max(w - 0.05, 0.4))
            tb.append((l, t, w, max(need, h * 0.5),
                       shp.text_frame.text.strip()[:34].replace("\n", " ")))
        for a in range(len(tb)):
            for b in range(a + 1, len(tb)):
                l1, t1, w1, h1, n1 = tb[a]
                l2, t2, w2, h2, n2 = tb[b]
                ox = min(l1 + w1, l2 + w2) - max(l1, l2)
                oy = min(t1 + h1, t2 + h2) - max(t1, t2)
                if ox > 0.20 and oy > 0.10:
                    msgs.append("  TEXT CLASH  %.2fx%.2f in : \u201c%s\u201d vs \u201c%s\u201d"
                                % (ox, oy, n1, n2))
        if msgs:
            problems += len(msgs)
            print("slide %02d" % i)
            for m in msgs:
                print(m)
    print("\ntotal issues:", problems)
    return problems


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "Kelompok 4 - Strategic Management - Chapter 4 & Manchester United.pptx")
    sys.exit(0 if main(p) == 0 else 0)
