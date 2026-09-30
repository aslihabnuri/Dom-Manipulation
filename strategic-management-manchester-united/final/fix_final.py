from pptx import Presentation
from pptx.util import Inches
import copy

SRC, DST = 'final.pptx', 'final_v2.pptx'
p = Presentation(SRC)
S = lambda n: p.slides[n - 1]
def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)

def set_para(shape_, idx, text):
    """Replace text of the idx-th NON-EMPTY paragraph, keeping its first run's formatting."""
    paras = [q for q in shape_.text_frame.paragraphs if q.runs]
    para = paras[idx]
    runs = para.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def set_text(shape_, text):
    """Single-paragraph replacement, keeping first run formatting."""
    tf = shape_.text_frame
    set_para(shape_, 0, text)
    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)

def notes_text(slide):
    return slide.notes_slide.notes_text_frame.text

def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def edit_notes(slide, old, new):
    t = notes_text(slide)
    assert t.count(old) == 1, (old[:40], t.count(old))
    set_notes(slide, t.replace(old, new))

# ---------- slide 3 / 18: consistent part labels ----------
set_text(shape(S(3), 3), "BAGIAN 1  ·  FRAMEWORK CHAPTER 4")
set_text(shape(S(18), 5), "BAGIAN 2  ·  STUDI KASUS")
edit_notes(S(2), "bagian A bisa dipercepat", "Bagian 1 bisa dipercepat")
edit_notes(S(3), "Ini peta jalan seluruh bagian A.", "Ini peta jalan seluruh Bagian 1.")
edit_notes(S(19), "benang merah Bagian B:", "benang merah Bagian 2:")

# ---------- slide 4: AirAsia is the group's example ----------
set_text(shape(S(4), 14), "CONTOH KAMI: AIRASIA")
edit_notes(S(4), "Mulai dengan: sebelum menilai strategi",
           "Catatan: AirAsia adalah contoh kelompok kami, bukan contoh dari buku. "
           "Sebutkan itu kalau ditanya.\n\nMulai dengan: sebelum menilai strategi")

# ---------- slide 18 notes: dilemma box moved to slide 19 ----------
edit_notes(S(18), "Dilema di kotak terakhir dikutip dari hlm. 588. Perhatikan kata kasusnya:",
           "Dilema Gill dibahas di slide berikutnya. Perhatikan kata kasusnya (hlm. 588):")

# ---------- slide 19: stock photo is not Old Trafford ----------
set_text(shape(S(19), 22), "Foto ilustrasi stadion, bukan Old Trafford")

# ---------- slide 21: Table 6.2 = total points, not performance points ----------
set_text(shape(S(21), 5), "Indikator kedua Thompson: apakah posisi bersaing MU membaik? "
                          "Data kompetisi di kasus menjawabnya.")
set_text(shape(S(21), 26), "TOTAL POIN EROPA 2000–09  ·  TABLE 6.2")
set_text(shape(S(21), 46), "Di antara enam klub ini, skuad MU paling besar (34 pemain) dan termuda "
                           "kedua (25,6 tahun) setelah Arsenal (Table 6.7). Cocok dengan strategi "
                           "memadukan pemain muda dan pemain berpengalaman (hlm. 583).")
set_notes(S(21),
    "Sumber angka: Table 6.2 (total poin Eropa), Table 6.1 (juara liga), hlm. 584 (Liga "
    "Champions), dan Table 6.7 untuk catatan ukuran dan usia skuad.\n\n"
    "Hati-hati dengan dua ukuran 'poin' di kasus, karena mudah tertukar:\n"
    "- Total poin 2000–09 (Table 6.2, yang ada di grafik slide ini): dihitung dari prestasi "
    "di liga domestik, piala domestik, Liga Champions, dan Intertoto Cup, disesuaikan dengan "
    "tingkat kesulitan liga. MU nomor satu, 1.460.\n"
    "- Poin performa 2000–09 (Table 6.7, dipakai di slide 23): ukuran yang disandingkan kasus "
    "dengan belanja transfer. Di sini MU nomor dua (192,5), tipis di bawah Barcelona (200,0).\n"
    "Kalau ditanya kenapa berbeda: dua tabel, dua metode; kasus tidak menjelaskan rumus poin "
    "performa.\n\n"
    "Soal kartu 2008: itu Liga Champions (European Cup) KEDUA di era Ferguson, setelah 1999. "
    "Jangan sebut 'gelar Eropa kedua', karena kasus hlm. 584 juga mencatat Piala Winners "
    "Eropa 1991.\n\n"
    "Soal ukuran skuad di bawah grafik: angkanya dari Table 6.7, dan klaimnya hanya berlaku "
    "untuk enam klub di grafik. Di Table 6.7 lengkap, Roma punya 46 pemain dan Lyon 34. Jadi "
    "jangan bilang skuad MU terbesar di Eropa. Strategi 'memadukan pemain muda dan "
    "berpengalaman' dikutip dari hlm. 583.")

# ---------- slide 22: subtitle wording ----------
set_text(shape(S(22), 5), "Indikator pertama Thompson: apakah kekuatan keuangan dan laba MU membaik? "
                          "Appendix kasus menjawabnya.")

# ---------- slide 23 notes: stray reminder ----------
edit_notes(S(23), "\n\nNotes: Grafik ditambahkan Chelsea", "")

# ---------- slide 30: benchmarking — honest club set, correct framing ----------
s30 = S(30)
set_text(shape(s30, 5), "Benchmarking untuk Q4: seberapa kompetitif struktur biaya MU? Ukurannya bukan "
                        "besar pendapatan, tapi berapa yang tersisa sebagai laba.")
set_text(shape(s30, 11), "MARGIN EBITDA, %  ·  10 KLUB BERPENDAPATAN TERBESAR DI TABLE 6.5")
rows = [("Man United", 31.4), ("Roma", 25.0), ("Barcelona", 22.3), ("Arsenal", 19.3),
        ("AC Milan", 17.6), ("Liverpool", 15.8), ("Real Madrid", 14.1), ("Bayern", 12.7),
        ("Inter", 9.9), ("Chelsea", -3.1)]
name_ids = [12, 15, 18, 21, 24, 27, 30, 33, 36, 39]
bar_ids = [13, 16, 19, 22, 25, 28, 31, 34, 37, 40]
val_ids = [14, 17, 20, 23, 26, 29, 32, 35, 38, 41]
AXIS, SCALE = 2.70, 3.50 / 31.4
for (nm, v), nid, bid, vid in zip(rows, name_ids, bar_ids, val_ids):
    set_text(shape(s30, nid), nm)
    set_text(shape(s30, vid), ("%.1f" % v).replace(".", ",").replace("-", "−"))
    bar, val = shape(s30, bid), shape(s30, vid)
    w = abs(v) * SCALE
    bar.width = Inches(w)
    if v >= 0:
        bar.left = Inches(AXIS)
        val.left = Inches(AXIS + w + 0.08)
    else:
        bar.left = Inches(AXIS - w)
        val.left = Inches(AXIS + 0.08)
set_text(shape(s30, 45), "Dua jalan mendapat pemain (hlm. 580)")
box = shape(s30, 46)
set_para(box, 0, "Membeli bintang: Real Madrid, Chelsea. Return on sales 0,4% dan −60,4% (Table 6.6).")
set_para(box, 1, "Mengembangkan sendiri: Bayern, Barcelona, Arsenal, Valencia. Bayern 4,7%, Arsenal 5,6%.")
set_para(box, 2, "MU memadukan keduanya: bakat akademi dipadukan pemain berpengalaman, lalu pemain "
                 "dijual saat klub lain menilai lebih tinggi (hlm. 583, 586). Return on sales 12,6%, "
                 "tertinggi di Table 6.6.")
set_text(shape(s30, 48), "BATAS KLAIM INI")
box = shape(s30, 49)
set_para(box, 0, "Margin MU tertinggi di antara klub besar. Dari 15 klub di Table 6.5, hanya Lyon "
                 "(38,5%) yang lebih tinggi, dengan pendapatan kurang dari separuh MU.")
set_para(box, 1, "Benchmarking menguji biaya dan hasil aktivitas (Q4). VRIN Test menguji resource (Q3). "
                 "Dua alat, dua pertanyaan berbeda.")
set_notes(s30,
    "Slide ini menjawab bagian kedua Q4: apakah struktur biaya MU kompetitif? Caranya "
    "benchmarking, membandingkan hasil akhir yang sama untuk semua klub: berapa persen "
    "pendapatan yang tersisa sebagai laba.\n\n"
    "Dua ukuran dipakai:\n"
    "- Margin EBITDA, dihitung kami dari Table 6.5 (EBITDA dibagi pendapatan). Grafik memuat "
    "10 klub berpendapatan terbesar di tabel itu. MU 31,4%, tertinggi di antara mereka.\n"
    "- Return on sales 2000–2006 dari Table 6.6, angka kasus sendiri. MU 12,6%, tertinggi dari "
    "10 klub.\n\n"
    "Batas klaim, sebutkan sendiri sebelum ditanya: Table 6.5 memuat 15 klub. Dari 15 itu, "
    "Lyon punya margin EBITDA 38,5% (€59,9 jt dari €155,7 jt), di atas MU, tapi pendapatannya "
    "kurang dari separuh MU. Tottenham 30,8%, Roma 25,0%. Jadi kalimat yang benar adalah "
    "'tertinggi di antara klub besar', bukan 'tertinggi di Eropa'.\n\n"
    "Kaitkan dengan dua jalan mendapat pemain (hlm. 580): membeli atau mengembangkan. Real "
    "Madrid dan Chelsea membeli, dan return on sales-nya 0,4% dan minus 60,4%. Bayern, "
    "Barcelona, Arsenal, dan Valencia disebut kasus menekankan pengembangan pemain sendiri; "
    "Bayern 4,7%, Arsenal 5,6%. MU memadukan keduanya (hlm. 583 dan 586): bakat akademi "
    "dipadukan dengan pemain berpengalaman, dan pemain dijual saat klub lain menilainya lebih "
    "tinggi. Hasilnya 12,6%.\n\n"
    "Jangan katakan benchmarking 'membuktikan VRIN'. Dua alat itu menjawab pertanyaan "
    "berbeda: VRIN (Q3) menguji apakah resource bisa ditiru atau digantikan pesaing; "
    "benchmarking (Q4) menguji apakah biaya dan hasil aktivitas kita lebih baik dari pesaing. "
    "Slide ini bukti untuk Q4.")

# ---------- slide 32: priority list wording + evidence ----------
s32 = S(32)
set_text(shape(s32, 18), "Bagaimana bersaing di bursa transfer tanpa dana pemilik?")
set_text(shape(s32, 19), "Utang £616 jt tidak menghambat belanja pemain ('proved groundless', hlm. 584), "
                         "tapi MU tak dapat dana pemilik seperti City yang belanja £185 jt (hlm. 579).")
set_text(shape(s32, 24), "Tidak ada lagi pemain MU di peringkat FIFA Table 6.3. Media menunggu apakah "
                         "kas Ronaldo dipakai membeli Ribéry (hlm. 588).")
set_text(shape(s32, 28), "Apakah perlu director of football untuk menopang sistem?")
set_text(shape(s32, 29), "Seluruh sistem pemanduan bakat, latihan, dan taktik dibangun satu orang "
                         "(hlm. 573). Di Inggris posisi ini sering memicu konflik dengan manajer (hlm. 581).")
set_text(shape(s32, 38), "Apakah pemilik tetap pasif setelah Ferguson pergi?")
set_text(shape(s32, 39), "Peran pasif Glazer berjalan karena Ferguson dan Gill diberi kebebasan (hlm. 584). "
                         "Di Chelsea, campur tangan pemilik memicu kepergian Mourinho (hlm. 581).")
set_notes(s32,
    "Ingatkan aturan Q6 dari slide 17: priority list berisi PERTANYAAN, bukan jawaban. Keenam "
    "isu di slide ini berbentuk 'bagaimana', 'apakah perlu', dan 'apakah', persis pola buku "
    "(hlm. 118). Jangan menjawabnya di slide ini; jawaban ada di tahap penyusunan strategi.\n\n"
    "Sumber tiap isu:\n"
    "1. Suksesi: hlm. 573 (Ferguson belum menyatakan niat pensiun; Gill memperkirakan akhir "
    "musim 2009–10) dan hlm. 588 (dilema orang dalam vs orang luar, dan pentingnya "
    "mempertahankan sistem latihan, pemanduan bakat, dan pengembangan tim).\n"
    "2. Pasar transfer tanpa dana pemilik: hlm. 579 (karena akuisisi dibiayai utang, MU tidak "
    "mendapat suntikan dana seperti Chelsea dan Manchester City, yang belanja £185 juta pada "
    "2008–09). PENTING: jangan bilang utang membatasi pembelian pemain. Hlm. 584 menyebut "
    "ketakutan itu 'proved groundless'.\n"
    "3. Skuad setelah Ronaldo: Table 6.3 (Ronaldo satu-satunya pemain MU di peringkat FIFA), "
    "Table 6.4 (dijual ke Real Madrid, rekor $93,1 juta), hlm. 588 (media menunggu apakah "
    "uangnya dipakai menawar Ribéry dari Bayern).\n"
    "4. Director of football: hlm. 573 (seluruh infrastruktur dibangun Ferguson) dan hlm. 581 "
    "(posisi ini lazim di Eropa daratan, tapi di Inggris sering memicu konflik dengan manajer, "
    "contohnya Chelsea dan Newcastle).\n"
    "5. Pendapatan komersial: hlm. 573, pandangan Gill bahwa pendapatan komersial mengalir "
    "selama tim menang dan bermain menarik.\n"
    "6. Sikap pemilik: hlm. 584 (Glazer memilih peran pasif; Ferguson dan Gill diberi kebebasan) "
    "dan hlm. 581 (di Chelsea, gesekan dengan Abramovich soal pembelian pemain dan pemilihan "
    "tim berujung pada kepergian Mourinho).\n\n"
    "Pembacaan di kotak bawah mengikuti diagnosis buku: kalau isunya berat, menyusun strategi "
    "baru harus jadi agenda utama. Isu 1 dan 4 berat karena menyangkut resource yang lolos "
    "VRIN tapi melekat pada satu orang. Isu lainnya bisa ditangani dengan menyempurnakan "
    "strategi yang ada.\n\n"
    "Kalau dosen bertanya 'jadi apa rekomendasi kalian?', jawab singkat, dan sebut bahwa ini "
    "sudah tahap penyusunan strategi, bukan Q6: (1) lembagakan sistem Ferguson menjadi prosedur "
    "klub sebelum ia pergi; (2) pilih penerus dari dalam dengan masa transisi, Ferguson "
    "mendampingi; (3) tolak opsi revolusi seperti Mourinho, karena kasus sendiri menyebut itu "
    "berarti membongkar sebagian besar infrastruktur Ferguson (hlm. 588); (4) lindungi mesin "
    "komersial yang melekat pada klub, terutama India dan Cina.")

# ---------- slide 33 notes: add the recommendation question ----------
t = notes_text(S(33))
set_notes(S(33), t.rstrip() + "\n\n"
    "7. Jadi apa rekomendasi kelompok untuk Gill?\n"
    "Ini sudah tahap penyusunan strategi, di luar Chapter 4, tapi siapkan jawabannya: "
    "lembagakan sistem Ferguson menjadi prosedur klub; pilih penerus dari dalam dengan masa "
    "transisi; tolak opsi revolusi karena membongkar resource yang lolos VRIN (hlm. 588); dan "
    "lindungi mesin komersial yang melekat pada klub.")

p.save(DST)
# python-pptx swaps the jpg Default for per-part Overrides; restore the original declaration style
import zipfile, shutil, re, os
tmp = DST + '.tmp'
with zipfile.ZipFile(DST) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == '[Content_Types].xml':
            ct = data.decode('utf-8')
            ct = re.sub(r'<Override PartName="/ppt/media/[^"]+\.jpg" ContentType="image/jpeg"/>', '', ct)
            if 'Extension="jpg"' not in ct:
                ct = ct.replace('<Default Extension="jpeg" ContentType="image/jpeg"/>',
                                '<Default Extension="jpeg" ContentType="image/jpeg"/><Default Extension="jpg" ContentType="image/jpeg"/>', 1)
            data = ct.encode('utf-8')
        zout.writestr(item, data)
os.replace(tmp, DST)
print("saved", DST)
